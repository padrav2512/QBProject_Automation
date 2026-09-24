import pytest
from docx import Document
from docx.oxml.ns import qn
from cms_processor import ProcessorOptions, process_docx
from safe_math import parse


@pytest.mark.parametrize('text', ['r² − d²', '√(r² + d²)', '2√(r² − d²)', '2(r − d)', 'c = 2√(r² − d²)', '1/2'])
def test_explicit_expressions(text):
    equation = parse(text)
    for node in equation.iter(qn('m:r')):
        value = node.find(qn('m:t')).text
        style = node.find(qn('m:rPr')).find(qn('m:sty')).get(qn('m:val'))
        assert style == ('i' if value.isalpha() else 'p')


@pytest.mark.parametrize('text', ['1/2r', '√r² − d²', 'AB', '√27^2', '1/2/3', 'r^', '√(r² − d²'])
def test_ambiguous_expressions_rejected(text):
    with pytest.raises(ValueError):
        parse(text)


def test_case_study_and_inline_expression(tmp_path):
    doc = Document()
    for text in ['@Question: 81@ Case study', 'Type: FIB', 'Question:',
                 'Reason (R): radius r, chord length c, and distance d are related by c = 2√(r² − d²).',
                 'Answers:', '1/2', 'Question 82', 'Question:', 'Choose.', 'Choices:',
                 'a) r² − d² units', 'b) √(r² + d²) units', 'c) 2√(r² − d²) units', 'd) 2(r − d) units', 'Answer: c']:
        doc.add_paragraph(text)
    source = tmp_path / 'source.docx'; doc.save(source)
    result = process_docx(source, source.name, ProcessorOptions(start_question_number=81), tmp_path / 'storage')
    out = Document(result.output_path)
    texts = [p.text for p in out.paragraphs]
    assert texts.count('@Solution:@') == 2
    assert 'Case study' in texts
    assert len(out.element.xpath('.//m:rad')) == 3
    assert len(out.element.xpath('.//m:f')) == 1
    reason = next(p for p in out.paragraphs if p.text.startswith('Reason'))
    assert reason.text == 'Reason (R): radius r, chord length c, and distance d are related by .'
    choices = [p for p in out.paragraphs if p.text.startswith(('@1@', '@2@', '@3@', '@4@'))]
    assert len(choices) == 4
    assert all('units' in p.text and p._p.xpath('.//m:oMath') for p in choices)
    assert '@correct answer@' in choices[2].text
