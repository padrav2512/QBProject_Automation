from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree
from PIL import Image
import pytest

from cms_processor import PROCESSOR_BUILD_ID, ProcessorOptions, process_docx


def add_radical_answer(paragraph):
    from cms_processor import _math_run, _math_radical
    equation = OxmlElement("m:oMath")
    equation.append(_math_run("2"))
    equation.append(_math_radical("13"))
    paragraph._p.append(equation)


def italicised_text(paragraph):
    return [(run.text, run.italic is True) for run in paragraph.runs if run.text]


def test_v12_authoring_rules(tmp_path: Path):
    from cms_processor import _math_fraction

    doc = Document()
    doc.add_paragraph("Question 1: Easy, Knowledge")
    doc.add_paragraph("Question type: FIB")
    question = doc.add_paragraph("Riya drank ")
    existing = OxmlElement("m:oMath")
    existing.append(_math_fraction("1", "3"))
    question._p.append(existing)
    question.add_run(" litres in the morning and 1/4 litres in the evening. What did she drink altogether ?")
    doc.add_paragraph("Correct answer: 7/12 litres")

    doc.add_paragraph("Question 2: Easy, Knowledge")
    doc.add_paragraph("Question type: FIB")
    doc.add_paragraph("Find the values of √144, √25 and √(r² + d²). Preserve √a+b, √5/9 and √r² for review.")
    doc.add_paragraph("Answer: x = √144")

    doc.add_paragraph("Question 3: Easy, Knowledge")
    doc.add_paragraph("Question type: MCQ")
    doc.add_paragraph("In the number 7³, what is the base?")
    for text in ["a) 4", "b) 7", "c) 11", "d) 28", "Correct answer: b) 7"]:
        doc.add_paragraph(text)

    doc.add_paragraph("Question 4: Easy, Comprehension")
    doc.add_paragraph("Question type: FIB")
    doc.add_paragraph("Point A lies on line BC. A circle has radius r. Let m be a number. The length is 5 m. Option A is shown.")
    doc.add_paragraph("Answer: 5 cm")

    source = tmp_path / "v12_rules.docx"
    doc.save(source)
    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    out = Document(result.output_path)
    texts = [paragraph.text for paragraph in out.paragraphs]

    assert PROCESSOR_BUILD_ID == "2026.09.25-processing-sections-v13"
    assert len(out.element.xpath(".//m:f")) >= 3
    assert "1/4" not in "\n".join(texts)
    assert "7/12" not in "\n".join(texts)
    assert any("What did she drink altogether?" in text for text in texts)
    assert len(out.element.xpath(".//m:rad")) >= 4
    assert any("√a + b" in text and "√5/9" in text and "√r²" in text for text in texts)
    assert "@Type: MCQ@" in texts
    assert "@2@ 7 @correct answer@" in texts
    assert "Question type: MCQ" not in texts
    assert not any(text.startswith("Correct answer:") for text in texts)

    geometry = next(paragraph for paragraph in out.paragraphs if paragraph.text.startswith("Point A"))
    styled = italicised_text(geometry)
    italic_text = "".join(text for text, italic in styled if italic)
    assert "A" in italic_text and "BC" in italic_text and "r" in italic_text and "m" in italic_text
    article_index = geometry.text.index("A circle")
    offset = 0
    article_run = None
    for run in geometry.runs:
        if offset <= article_index < offset + len(run.text):
            article_run = run
            break
        offset += len(run.text)
    assert article_run is not None and article_run.italic is not True
    assert any(run.text == "m" and run.italic is True for run in geometry.runs)
    unit_index = geometry.text.index("5 m") + 2
    offset = 0
    unit_run = None
    for run in geometry.runs:
        if offset <= unit_index < offset + len(run.text):
            unit_run = run
            break
        offset += len(run.text)
    assert unit_run is not None and unit_run.italic is not True


def test_v12_rupee_conversion_and_conflicting_mcq_key(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Question 1")
    doc.add_paragraph("Question type: FIB")
    doc.add_paragraph("A book costs ₹50, a pen costs ₹ 50 and a bag costs ₹1,250.50.")
    doc.add_paragraph("Answer: ₹1,350.50")
    doc.add_paragraph("Question 2")
    doc.add_paragraph("Question type: MCQ")
    doc.add_paragraph("Choose the number.")
    for text in ["a) 4", "b) 7", "c) 11", "d) 28", "Correct answer: b) 9"]:
        doc.add_paragraph(text)
    source = tmp_path / "currency_conflict.docx"
    doc.save(source)

    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    out = Document(result.output_path)
    combined = "\n".join(paragraph.text for paragraph in out.paragraphs)
    assert "₹" not in combined
    assert "Rs 50" in combined and "Rs 1,250.50" in combined and "Rs 1,350.50" in combined
    assert any(f.status == "manual_review" and "does not match" in f.message and f.question == 2 for f in result.findings)


def test_v12_inline_equation_spacing_superscript_and_variable_style(tmp_path: Path):
    from cms_processor import _math_fraction

    doc = Document()
    doc.add_paragraph("Question 1")
    doc.add_paragraph("Question type: FIB")
    doc.add_paragraph("Find the roots of the equation x²−25=0.")
    doc.add_paragraph("Answer: 5, -5")

    doc.add_paragraph("Question 2")
    doc.add_paragraph("Question type: FIB")
    fraction_question = doc.add_paragraph("Riya drank ")
    existing_fraction = OxmlElement("m:oMath")
    existing_fraction.append(_math_fraction("1", "3"))
    fraction_question._p.append(existing_fraction)
    fraction_question.add_run(" litres in the morning and 1/4 litres in the evening.")
    doc.add_paragraph("Answer: 7/12 litres")

    doc.add_paragraph("Question 3")
    doc.add_paragraph("Question type: FIB")
    radical = doc.add_paragraph()
    radical.add_run("Find the value of √(144-r")
    exponent = radical.add_run("3")
    exponent.font.superscript = True
    radical.add_run(").")
    doc.add_paragraph("Answer: 12")

    source = tmp_path / "inline_math_boundaries.docx"
    doc.save(source)
    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    out = Document(result.output_path)

    equation_paragraph = next(p for p in out.paragraphs if "Find the roots" in p.text)
    x_styles = [
        node.get(qn("m:val"))
        for node in equation_paragraph._p.xpath(".//m:r[m:t='x']/m:rPr/m:sty")
    ]
    assert x_styles == ["i"]

    fraction_paragraph = next(p for p in out.paragraphs if "Riya drank" in p.text)
    fraction_equations = fraction_paragraph._p.xpath(".//m:oMath[m:f]")
    assert len(fraction_equations) == 2
    for equation in fraction_equations:
        direct_text = [
            run.find(qn("m:t")).text
            for run in equation.findall(qn("m:r"))
            if run.find(qn("m:t")) is not None
        ]
        assert direct_text[0] == " " and direct_text[-1] == " "

    radical_paragraph = next(p for p in out.paragraphs if "Find the value of" in p.text)
    assert radical_paragraph._p.xpath(".//m:rad//m:sSup[m:e//m:t='r'][m:sup//m:t='3']")
    assert not any(
        finding.status == "manual_review" and finding.question == 3 and "radical" in finding.message.lower()
        for finding in result.findings
    )


@pytest.mark.parametrize("all_bold", [False, True])
def test_author_formatting_and_whole_choice_bold(tmp_path: Path, all_bold):
    doc = Document()
    doc.add_paragraph("Question 1")
    doc.add_paragraph("A circle has radius 15 m. Find the distance to line ")
    doc.paragraphs[-1].add_run("AB").italic = True
    for i, label in enumerate("abcd"):
        p = doc.add_paragraph(label + ") ")
        r = p.add_run("AB" if i == 0 else "15 m")
        r.italic = i == 0
        r.bold = all_bold or i == 0
    doc.add_paragraph("Answer: a")
    doc.add_paragraph("Solution:")
    p = doc.add_paragraph()
    p.add_run("Use ").bold = True
    r = p.add_run("AB")
    r.bold = True
    r.italic = True
    doc.add_paragraph("Question 2")
    doc.add_paragraph("Type: FIB")
    doc.add_paragraph("Find the length.")
    p = doc.add_paragraph("Answer: ")
    r = p.add_run("AB")
    r.bold = True
    r.italic = True
    source = tmp_path / "authored.docx"
    doc.save(source)
    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    out = Document(result.output_path)
    question = next(p for p in out.paragraphs if p.text.startswith("A circle"))
    assert all(r.italic is not True for r in question.runs if r.text != "AB")
    assert next(r for r in question.runs if r.text == "AB").italic is True
    choice = next(p for p in out.paragraphs if p.text.startswith("@1@"))
    r = next(r for r in choice.runs if r.text == "AB")
    assert r.italic is True
    assert r.bold is (not all_bold)
    solution = next(p for p in out.paragraphs if p.text == "Use AB")
    assert all(r.bold is False for r in solution.runs if r.text)
    assert next(r for r in solution.runs if r.text == "AB").italic is True
    answer = next(p for p in out.paragraphs if p.text == "AB")
    r = next(r for r in answer.runs if r.text == "AB")
    assert r.bold is True and r.italic is True


@pytest.mark.parametrize("operator", ["f", "sSup", "sSub"])
def test_fraction_and_script_answers_are_preserved(tmp_path: Path, operator):
    containers = {"f": ("num", "den"), "sSup": ("e", "sup"), "sSub": ("e", "sub")}
    left, right = containers[operator]
    equation = etree.fromstring((f'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:{operator}><m:{left}><m:r><m:t>3</m:t></m:r></m:{left}><m:{right}><m:r><m:t>2</m:t></m:r></m:{right}></m:{operator}></m:oMath>').encode())
    doc = Document()
    doc.add_paragraph("Question 1")
    doc.add_paragraph("Find the value.")
    p = doc.add_paragraph("Answer: ")
    p._p.append(equation)
    source = tmp_path / "math.docx"
    doc.save(source)
    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    out = Document(result.output_path)
    assert len(out.element.xpath(f".//m:{operator}")) == 1
    assert [n.text for n in out.element.xpath(".//m:t")] == ["3", "2"]


def test_equation_loss_blocks_output(tmp_path: Path, monkeypatch):
    import cms_processor
    doc = Document()
    doc.add_paragraph("Question 1")
    doc.add_paragraph("Answers:")
    add_radical_answer(doc.add_paragraph())
    source = tmp_path / "protected.docx"
    doc.save(source)
    def lose_equations(document, findings):
        for node in document.element.xpath(".//m:oMath"):
            node.getparent().remove(node)
    monkeypatch.setattr(cms_processor, "_repair_spacing", lose_equations)
    with pytest.raises(ValueError, match="original Word equation"):
        process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    assert not list((tmp_path / "storage" / "processed").glob("*.docx"))


@pytest.mark.parametrize("units", [" cm", ""])
def test_native_equation_fib_answer_survives(tmp_path: Path, units):
    doc = Document()
    doc.add_paragraph("Question 37: Average, Comprehension")
    doc.add_paragraph("A circle has radius 7 cm, and a chord is 6 cm from its centre. Find its length.")
    answer = doc.add_paragraph("Answer: ")
    add_radical_answer(answer)
    answer.add_run(units)
    source = tmp_path / "radical.docx"
    doc.save(source)
    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    out = Document(result.output_path)
    texts = [p.text for p in out.paragraphs]
    answer_out = out.paragraphs[texts.index("@Answers:@") + 1]
    assert [n.text for n in answer_out._p.xpath(".//m:t") if (n.text or "").strip()] == ["2", "13"]
    assert len(answer_out._p.xpath(".//m:rad")) == 1
    assert answer_out.text.strip() == units.strip()


def test_native_equation_choices_and_subparts_survive(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Question 1")
    part = doc.add_paragraph("(i) Choose ")
    add_radical_answer(part)
    for label in "abcd":
        p = doc.add_paragraph(label + ") ")
        add_radical_answer(p)
    doc.add_paragraph("Answer: c")
    doc.add_paragraph("Solution:")
    add_radical_answer(doc.add_paragraph("The length is "))
    source = tmp_path / "choices.docx"
    doc.save(source)
    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    out = Document(result.output_path)
    assert len(out.element.xpath(".//m:rad")) == 6
    correct = next(p for p in out.paragraphs if "@correct answer@" in p.text)
    assert correct.text.startswith("@3@")
    assert correct._p.xpath(".//m:rad")


@pytest.mark.parametrize("headings", [False, True])
def test_plain_mcq_with_solution_preserves_sections(tmp_path: Path, headings):
    doc = Document()
    lines = ["Question 26: Average, Comprehension"]
    if headings:
        lines += ["Question:"]
    lines += ["Let M be the midpoint of chord AB in a circle with centre O.",
              "Which criterion proves triangles OMA and OMB congruent?"]
    if headings:
        lines += ["Options:"]
    lines += ["a) ASA", "b) SAS", "c) SSS", "d) RHS", "Answer: c",
              "Solution:", "OA = OB (radii), AM = BM (midpoint), OM common."]
    for line in lines:
        assert "@" not in line
        doc.add_paragraph(line)
    source = tmp_path / "plain.docx"
    doc.save(source)
    result = process_docx(source, source.name, ProcessorOptions(start_question_number=26), tmp_path / "storage")
    texts = [p.text for p in Document(result.output_path).paragraphs]
    assert "@Type: MCQ@" in texts
    assert texts.count("@Choices:@") == 1
    assert texts.count("@Solution:@") == 1
    assert "@3@ SSS @correct answer@" in texts
    assert sum("@correct answer@" in t for t in texts) == 1
    assert "Answer: c" not in texts
    assert "@Objective: Comprehension @" in texts
    solution = texts.index("@Solution:@")
    assert texts[solution + 1] == lines[-1]
    assert texts[solution + 2] == "@e@"
    assert texts[texts.index("@Choices:@") - 1] == "@e@"


def test_plain_fib_with_solution_and_following_mcq(tmp_path: Path):
    doc = Document()
    for line in ["Question 1", "Type: FIB", "Question:", "How many?", "Answer: 42",
                 "Solution:", "Six groups of seven.", "Question 2", "Which number?",
                 "a) 1", "b) 2", "c) 3", "d) 4", "Answer: b", "Solution:", "Two is correct."]:
        doc.add_paragraph(line)
    source = tmp_path / "mixed.docx"
    doc.save(source)
    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    texts = [p.text for p in Document(result.output_path).paragraphs]
    assert result.question_count == 2
    assert texts.count("@Solution:@") == 2
    assert "@Type: FIB@" in texts and "@Type: MCQ@" in texts
    assert texts[texts.index("@Answers:@") + 1] == "42"
    assert "Six groups of seven." in texts
    assert "@2@ 2 @correct answer@" in texts


def make_sample(path: Path) -> None:
    picture = path.with_name("sample_image.png")
    Image.new("RGB", (40, 30), "white").save(picture)
    doc = Document()
    doc.add_paragraph("Question 1")
    doc.add_paragraph("Type: FIB")
    doc.add_paragraph("Question:")
    p = doc.add_paragraph()
    run = p.add_run("Find  20=2×2×5,then write 3/10.")
    run.bold = True
    doc.add_picture(str(picture))
    doc.add_paragraph("Answers:")
    doc.add_paragraph("(a) 20  (b) 3/10")
    doc.add_paragraph("Solution:")
    s = doc.add_paragraph()
    sr = s.add_run("20=2×2×5")
    sr.bold = True
    doc.add_picture(str(picture))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Number"
    table.cell(0, 1).text = "Factor"
    table.cell(1, 0).text = "20"
    table.cell(1, 1).text = "5"
    doc.save(path)


def test_processing_assigns_tags_and_repairs_safe_formatting(tmp_path: Path):
    source = tmp_path / "sample.docx"
    make_sample(source)
    options = ProcessorOptions(
        project_question_prefix="project10436_q",
        start_question_number=7,
        start_snippet_id=220000,
    )
    result = process_docx(source, source.name, options, tmp_path / "storage")
    out = Document(result.output_path)
    texts = [p.text for p in out.paragraphs]

    assert "@Question: 7@" in texts
    assert "@Question id: project10436_q7 @" in texts
    assert "@New snippet id: 220000 @" in texts
    assert "(a) 20" in texts
    assert "(b) " in texts
    assert result.report_path.exists()
    assert result.manifest_path.exists()
    assert result.images_zip_path.exists()
    assert result.image_manifest_path.exists()
    assert result.image_count == 2
    assert any(f.rule == 7 and f.status == "fixed" for f in result.findings)

    question_index = texts.index("@Question:@")
    answer_index = texts.index("@Answers:@")
    for paragraph in out.paragraphs[question_index + 1 : answer_index]:
        assert all(run.bold is not True for run in paragraph.runs)

    assert all(run.bold is not True for row in out.tables[0].rows for cell in row.cells for p in cell.paragraphs for run in p.runs)

    xml = etree.fromstring(ZipFile(result.output_path).read("word/document.xml"))
    ns = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
    assert len(xml.xpath(".//m:oMath", namespaces=ns)) >= 2
    assert len(xml.xpath(".//m:f", namespaces=ns)) >= 1
    alt_texts = xml.xpath(".//wp:docPr/@descr", namespaces={"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"})
    assert "project10436_q7_1.gif" in alt_texts
    assert "project10436_a7.gif" in alt_texts


def test_selective_bold_emphasis_is_preserved(tmp_path: Path):
    source = tmp_path / "emphasis.docx"
    doc = Document()
    doc.add_paragraph("Question 1")
    doc.add_paragraph("Question:")
    question = doc.add_paragraph()
    question.add_run("Read the ")
    emphasis = question.add_run("important")
    emphasis.bold = True
    question.add_run(" word.")
    doc.add_paragraph("Answers:")
    doc.add_paragraph("10")
    doc.add_paragraph("Solution:")
    solution = doc.add_paragraph()
    solution.add_run("Use the ")
    solution_emphasis = solution.add_run("factor")
    solution_emphasis.bold = True
    solution.add_run(" shown.")
    doc.save(source)

    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    output = Document(result.output_path)
    question_out = next(p for p in output.paragraphs if p.text == "Read the important word.")
    solution_out = next(p for p in output.paragraphs if p.text == "Use the factor shown.")
    assert next(r for r in question_out.runs if r.text == "important").bold is True
    assert next(r for r in solution_out.runs if r.text == "factor").bold is True


def test_cms_tags_are_bold_and_author_italics_are_preserved(tmp_path: Path):
    source = tmp_path / "choice_variables.docx"
    doc = Document()
    for text in [
        "Question 1",
        "Type: MCQ",
        "Question:",
        "Choose the correct value.",
        "Choices:",
        "a) x = 4",
        "b) x = 5",
        "c) x = 6",
        "d) x = 7 @correct answer@",
        "Solution:",
        "The value of x is 7.",
    ]:
        doc.add_paragraph(text)
    doc.save(source)

    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    output = Document(result.output_path)

    bold_tags = {
        "@Question: 1@",
        "@Type: MCQ@",
        "@Question id: project10436_q1 @",
        "@New snippet id: 218989 @",
        "@Difficulty level: Average @",
        "@Objective: Application @",
        "@Question:@",
        "@Choices:@",
        "@Solution:@",
    }
    for paragraph in output.paragraphs:
        if paragraph.text in bold_tags:
            assert paragraph.runs
            assert all(run.bold is True for run in paragraph.runs if run.text)

    choice = next(p for p in output.paragraphs if p.text.startswith("@1@"))
    assert [n.text for n in choice._p.xpath(".//m:t")] == ["x", "=", "4"]
    assert all(run.italic is not True for run in choice.runs)

    solution = next(p for p in output.paragraphs if p.text == "The value of x is 7.")
    assert any(run.text == "x" and run.italic is True for run in solution.runs)

    end_tags = [p for p in output.paragraphs if p.text == "@e@"]
    assert end_tags
    assert all(all(run.bold is not True for run in p.runs) for p in end_tags)


def test_heading_with_difficulty_and_untagged_mcq_are_structured(tmp_path: Path):
    source = tmp_path / "untagged_mcq.docx"
    doc = Document()
    doc.add_paragraph("Q1) Easy")
    doc.add_paragraph("Which pair has a sum of 55?")
    doc.add_paragraph("a) 30 and 25")
    doc.add_paragraph("b) 35 and 20")
    doc.add_paragraph("20 and 35")
    doc.add_paragraph("25 and 30")
    doc.add_paragraph("Answer: b")
    doc.save(source)

    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    texts = [p.text for p in Document(result.output_path).paragraphs]

    assert result.question_count == 1
    assert "@Question: 1@" in texts
    assert "@Type: MCQ@" in texts
    assert "@Difficulty level: Easy @" in texts
    assert "@Question:@" in texts
    assert "@Choices:@" in texts
    assert "@2@ 35 and 20 @correct answer@" in texts
    assert "@Solution:@" in texts
    solution_index = texts.index("@Solution:@")
    assert texts[solution_index + 1] == "@e@"
    assert "The correct answer is option (b)." not in texts


def test_untagged_fib_answer_is_structured(tmp_path: Path):
    source = tmp_path / "untagged_fib.docx"
    doc = Document()
    doc.add_paragraph("Q6) Medium")
    doc.add_paragraph("A coin finishes at position P. Find P.")
    doc.add_paragraph("Answer 50")
    doc.save(source)

    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    texts = [p.text for p in Document(result.output_path).paragraphs]

    assert result.question_count == 1
    assert "@Type: FIB@" in texts
    assert "@Difficulty level: Average @" in texts
    assert "@Answers:@" in texts
    assert "50" in texts
    assert "@Solution:@" in texts
    solution_index = texts.index("@Solution:@")
    assert texts[solution_index + 1] == "@e@"
    assert "The answer is 50." not in texts


def test_tagged_question_suffix_is_counted_preserved_and_flagged(tmp_path: Path):
    source = tmp_path / "case_study_suffix.docx"
    doc = Document()
    for text in [
        "@Question: 29@ case study",
        "@Type: FIB@",
        "@Question id: project10436_q29 @",
        "@New snippet id: 219089 @",
        "@Difficulty level: Challenging @",
        "@Objective: Analysis @",
        "@Question:@",
        "Read the following information.",
        "@e@",
        "@Answers:@",
        "10",
        "@e@",
        "@Solution:@",
        "The answer is 10.",
        "@e@",
    ]:
        doc.add_paragraph(text)
    doc.save(source)

    options = ProcessorOptions(start_question_number=29, start_snippet_id=219089)
    result = process_docx(source, source.name, options, tmp_path / "storage")
    output = Document(result.output_path)
    texts = [p.text for p in output.paragraphs]

    assert result.question_count == 1
    assert "@Question: 29@" in texts
    question_marker_index = texts.index("@Question:@")
    assert texts[question_marker_index + 1] == "case study"
    assert texts.count("case study") == 1
    assert any(f.status == "manual_review" and "text after its closing tag" in f.message for f in result.findings)

    image_only = process_docx(
        source,
        source.name,
        ProcessorOptions(processing_mode="images_only", project_question_prefix="project10436_q"),
        tmp_path / "image_only_storage",
    )
    assert image_only.question_count == 1
    assert any(f.status == "manual_review" and "counted as a question" in f.message for f in image_only.findings)


def test_question_tags_accept_spaces_before_closing_at_sign(tmp_path: Path):
    source = tmp_path / "spaced_question_tags.docx"
    doc = Document()
    headings = ["@Question:10 @ case study", "@Question:24 @", "@Question: 42@"]
    for number, heading in zip((10, 24, 42), headings):
        for text in [
            heading,
            "@Type: FIB@",
            f"@Question id: project10434_q{number}@",
            f"@New snippet id: {220000 + number}@",
            "@Question:@",
            f"Question {number} text.",
            "@e@",
            "@Answers:@",
            str(number),
            "@e@",
            "@Solution:@",
            f"The answer is {number}.",
            "@e@",
        ]:
            doc.add_paragraph(text)
    doc.save(source)

    image_only = process_docx(
        source,
        source.name,
        ProcessorOptions(processing_mode="images_only", project_question_prefix="project10434_q"),
        tmp_path / "image_only_storage",
    )
    assert image_only.question_count == 3

    full = process_docx(
        source,
        source.name,
        ProcessorOptions(start_question_number=10, start_snippet_id=220010),
        tmp_path / "full_storage",
    )
    output_texts = [p.text for p in Document(full.output_path).paragraphs]
    assert full.question_count == 3
    assert "@Question: 10@" in output_texts
    assert "@Question: 11@" in output_texts
    assert "@Question: 12@" in output_texts
    assert "case study" in output_texts


def test_plain_heading_accepts_difficulty_and_objective(tmp_path: Path):
    source = tmp_path / "difficulty_objective_heading.docx"
    doc = Document()
    for text in [
        "Question 44: Easy, Comprehension",
        "Which number is even?",
        "a) 3",
        "b) 4",
        "c) 5",
        "d) 7",
        "Answer: b",
    ]:
        doc.add_paragraph(text)
    doc.save(source)

    result = process_docx(
        source,
        source.name,
        ProcessorOptions(project_question_prefix="projecttest_q", start_question_number=44, start_snippet_id=1044),
        tmp_path / "storage",
    )
    texts = [paragraph.text for paragraph in Document(result.output_path).paragraphs]

    assert result.question_count == 1
    assert "@Question: 44@" in texts
    assert "@Difficulty level: Easy @" in texts
    assert "@Objective: Comprehension @" in texts
    assert "@Type: MCQ@" in texts


def test_mcq_answer_key_with_closing_parenthesis_is_inferred_as_mcq(tmp_path: Path):
    source = tmp_path / "parenthesized_answer_key.docx"
    doc = Document()
    for text in [
        "Question 4: Easy, Comprehension",
        "Which statement is correct?",
        "a) First",
        "b) Second",
        "c) Third",
        "d) Fourth",
        "Answer: b)",
    ]:
        doc.add_paragraph(text)
    doc.save(source)

    result = process_docx(source, source.name, ProcessorOptions(), tmp_path / "storage")
    texts = [paragraph.text for paragraph in Document(result.output_path).paragraphs]

    assert result.question_count == 1
    assert "@Type: MCQ@" in texts
    assert "@Choices:@" in texts
    assert "@2@ Second @correct answer@" in texts
    assert "@Answers:@" not in texts


def test_safe_math_image_mode_converts_clear_choices_and_flags_ambiguous_scope(tmp_path: Path):
    source = tmp_path / "safe_math_choices.docx"
    doc = Document()
    for text in [
        "@Question: 5@",
        "@Type: MCQ@",
        "@Question id: project10434_q5@",
        "@New snippet id: 219065@",
        "@Difficulty level: Average@",
        "@Objective: Comprehension@",
        "@Question:@",
        "Choose the correct value.",
        "@e@",
        "@Choices:@",
        "@1@ x = √27 @correct answer@",
        "@2@ x = √36",
        "@3@ √3, √5/9, 1/√9",
        "@4@ x = 1/8",
        "@e@",
        "@Solution:@",
        "0.1666…, √2, 1/5",
        "x = -3^12",
        "@e@",
    ]:
        doc.add_paragraph(text)
    doc.save(source)

    result = process_docx(
        source,
        source.name,
        ProcessorOptions(processing_mode="images_math", project_question_prefix="project10434_q"),
        tmp_path / "safe_math_storage",
    )
    output = Document(result.output_path)
    output_texts = [p.text for p in output.paragraphs]
    assert "@Question id: project10434_q5@" in output_texts
    assert "@New snippet id: 219065@" in output_texts

    xml = etree.fromstring(ZipFile(result.output_path).read("word/document.xml"))
    ns = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
    assert len(xml.xpath(".//m:rad", namespaces=ns)) == 4
    assert len(xml.xpath(".//m:f", namespaces=ns)) == 2
    assert "@correct answer@" in "".join(xml.itertext())
    assert any(text.startswith("@3@") and "√5/9" in text and "1/√9" in text for text in output_texts)
    ambiguous = [f for f in result.findings if f.status == "manual_review" and "ambiguous scope" in f.message and f.question == 5]
    assert len(ambiguous) == 2
    assert not any("0.1666" in f.message and f.status == "manual_review" for f in result.findings)

    unchanged = process_docx(
        source,
        source.name,
        ProcessorOptions(processing_mode="images_only", project_question_prefix="project10434_q"),
        tmp_path / "images_only_storage",
    )
    unchanged_xml = etree.fromstring(ZipFile(unchanged.output_path).read("word/document.xml"))
    assert len(unchanged_xml.xpath(".//m:oMath", namespaces=ns)) == 0


def test_images_only_mode_preserves_existing_tags_and_text(tmp_path: Path):
    source = tmp_path / "already_tagged.docx"
    picture = tmp_path / "tagged_image.png"
    Image.new("RGB", (40, 30), "white").save(picture)
    doc = Document()
    for text in [
        "@Question: 7@",
        "@Type: MCQ@",
        "@Question id: existing_q7 @",
        "@New snippet id: 900 @",
        "@Difficulty level: Easy @",
        "@Objective: Application @",
        "@Question:@",
        "Choose the diagram.",
    ]:
        doc.add_paragraph(text)
    doc.add_picture(str(picture))
    for text in ["@e@", "@Choices:@", "@1@ One @correct answer@", "@2@ Two", "@3@ Three", "@4@ Four", "@e@", "@Solution:@", "One is correct.", "@e@"]:
        doc.add_paragraph(text)
    doc.save(source)
    original_texts = [p.text for p in doc.paragraphs]

    options = ProcessorOptions(processing_mode="images_only", project_question_prefix="project10436_q")
    result = process_docx(source, source.name, options, tmp_path / "storage")
    output = Document(result.output_path)

    assert result.question_count == 1
    assert [p.text for p in output.paragraphs] == original_texts
    xml = etree.fromstring(ZipFile(result.output_path).read("word/document.xml"))
    alt_texts = xml.xpath(".//wp:docPr/@descr", namespaces={"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"})
    assert "existing_q7_1.gif" in alt_texts


def test_exact_duplicate_question_images_reuse_first_filename(tmp_path: Path):
    source = tmp_path / "duplicate_images.docx"
    picture = tmp_path / "fraction_wall.png"
    Image.new("RGB", (60, 40), "white").save(picture)
    doc = Document()
    for question in (2, 5):
        for text in [f"@Question: {question}@", "@Type: FIB@", "@Question:@", "Use the fraction wall."]:
            doc.add_paragraph(text)
        doc.add_picture(str(picture))
        for text in ["@e@", "@Answers:@", "1/2", "@e@", "@Solution:@", "Use the wall.", "@e@"]:
            doc.add_paragraph(text)
    doc.save(source)

    options = ProcessorOptions(processing_mode="images_only", project_question_prefix="project1000_q")
    result = process_docx(source, source.name, options, tmp_path / "storage")
    xml = etree.fromstring(ZipFile(result.output_path).read("word/document.xml"))
    alt_texts = xml.xpath(".//wp:docPr/@descr", namespaces={"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"})

    assert alt_texts == ["project1000_q2_1.gif", "project1000_q2_1.gif"]
    assert result.image_count == 2
    with ZipFile(result.images_zip_path) as archive:
        gif_names = [name for name in archive.namelist() if name.lower().endswith(".gif")]
    assert gif_names == ["project1000_q2_1.gif"]
    duplicate_artifacts = [item for item in __import__('json').loads(result.image_manifest_path.read_text(encoding="utf-8"))["images"] if item["duplicate_of"]]
    assert len(duplicate_artifacts) == 1
    assert duplicate_artifacts[0]["duplicate_of"] == "project1000_q2_1.gif"


def test_images_only_mode_uses_project_id_from_existing_question_ids(tmp_path: Path):
    source = tmp_path / "existing_project_id.docx"
    picture = tmp_path / "diagram.png"
    Image.new("RGB", (40, 30), "white").save(picture)
    doc = Document()
    for text in [
        "@Question: 3@",
        "@Type: FIB@",
        "@Question id: project777_q3 @",
        "@New snippet id: 445566 @",
        "@Question:@",
        "Use the diagram.",
    ]:
        doc.add_paragraph(text)
    doc.add_picture(str(picture))
    for text in ["@e@", "@Answers:@", "3", "@e@", "@Solution:@", "Three.", "@e@"]:
        doc.add_paragraph(text)
    doc.save(source)

    options = ProcessorOptions(processing_mode="images_only", project_question_prefix="fallbackproject_q", start_snippet_id=1)
    result = process_docx(source, source.name, options, tmp_path / "storage")
    output = Document(result.output_path)
    output_texts = [p.text for p in output.paragraphs]
    xml = etree.fromstring(ZipFile(result.output_path).read("word/document.xml"))
    alt_texts = xml.xpath(".//wp:docPr/@descr", namespaces={"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"})

    assert "@Question id: project777_q3 @" in output_texts
    assert "@New snippet id: 445566 @" in output_texts
    assert alt_texts == ["project777_q3_1.gif"]


def test_safe_math_mode_preserves_existing_native_equations_images_and_ids(tmp_path: Path):
    source = tmp_path / "safe_math_preservation.docx"
    picture = tmp_path / "graph.png"
    Image.new("RGB", (40, 30), "white").save(picture)
    doc = Document()
    for text in [
        "@Question: 4@",
        "@Type: MCQ@",
        "@Question id: project10434_q4@",
        "@New snippet id: 219064@",
        "@Question:@",
        "Choose the correct value.",
        "@e@",
        "@Choices:@",
    ]:
        doc.add_paragraph(text)
    native_choice = doc.add_paragraph()
    native_choice.add_run("@1@ ")
    native_choice._p.append(
        etree.fromstring(
            b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x = 1</m:t></m:r></m:oMath>'
        )
    )
    native_choice.add_run(" @correct answer@")
    doc.add_paragraph("@2@ x = \u221a27")
    doc.add_paragraph("@3@ Three")
    doc.add_paragraph("@4@ Four")
    for text in ["@e@", "@Solution:@", "Use the graph."]:
        doc.add_paragraph(text)
    doc.add_picture(str(picture))
    doc.add_picture(str(picture))
    doc.add_paragraph("@e@")
    doc.save(source)

    result = process_docx(
        source,
        source.name,
        ProcessorOptions(processing_mode="images_math", project_question_prefix="project_q"),
        tmp_path / "storage",
    )
    output = Document(result.output_path)
    output_texts = [p.text for p in output.paragraphs]
    xml = etree.fromstring(ZipFile(result.output_path).read("word/document.xml"))
    ns = {
        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    }

    assert "@Question id: project10434_q4@" in output_texts
    assert "@New snippet id: 219064@" in output_texts
    assert len(xml.xpath(".//m:oMath", namespaces=ns)) == 2
    assert "@correct answer@" in "".join(xml.itertext())
    assert len(xml.xpath(".//wp:docPr", namespaces=ns)) == 2
    assert result.image_count == 2
    report = __import__("json").loads(result.report_path.read_text(encoding="utf-8"))
    assert report["processor_build_id"] == PROCESSOR_BUILD_ID


def test_question_type_is_corrected_only_when_structure_is_unambiguous(tmp_path: Path):
    source = tmp_path / "type_mismatches.docx"
    doc = Document()
    for text in [
        "@Question: 1@", "@Type: FIB@", "@Question id: project1_q1@", "@New snippet id: 10@",
        "@Question:@", "Choose one.", "@e@", "@Choices:@",
        "@1@ One @correct answer@", "@2@ Two", "@3@ Three", "@4@ Four", "@e@", "@Solution:@", "One.", "@e@",
        "@Question: 2@", "@Type: MCQ@", "@Question id: project1_q2@", "@New snippet id: 11@",
        "@Question:@", "Fill the blank.", "@e@", "@Answers:@", "5", "@e@", "@Solution:@", "Five.", "@e@",
        "@Question: 3@", "@Type: MCQ@", "@Question id: project1_q3@", "@New snippet id: 12@",
        "@Question:@", "Ambiguous structure.", "@e@", "@Choices:@",
        "@1@ One @correct answer@", "@2@ Two", "@3@ Three", "@4@ Four", "@e@", "@Answers:@", "1", "@e@", "@Solution:@", "One.", "@e@",
    ]:
        doc.add_paragraph(text)
    doc.save(source)

    repaired = process_docx(
        source,
        source.name,
        ProcessorOptions(processing_mode="images_math", project_question_prefix="project1_q"),
        tmp_path / "repair_storage",
    )
    repaired_texts = [paragraph.text for paragraph in Document(repaired.output_path).paragraphs]
    assert repaired_texts.count("@Type: MCQ@") == 2
    assert repaired_texts.count("@Type: FIB@") == 1
    fixed = [finding for finding in repaired.findings if finding.rule == 12 and finding.status == "fixed"]
    assert [(finding.question, "MCQ" in finding.message, "FIB" in finding.message) for finding in fixed] == [
        (1, True, True),
        (2, True, True),
    ]
    ambiguous = [finding for finding in repaired.findings if finding.rule == 12 and finding.status == "manual_review"]
    assert len(ambiguous) == 1 and ambiguous[0].question == 3

    preserved = process_docx(
        source,
        source.name,
        ProcessorOptions(processing_mode="images_only", project_question_prefix="project1_q"),
        tmp_path / "preserve_storage",
    )
    preserved_texts = [paragraph.text for paragraph in Document(preserved.output_path).paragraphs]
    assert preserved_texts.count("@Type: FIB@") == 1
    assert preserved_texts.count("@Type: MCQ@") == 2
    mismatches = [finding for finding in preserved.findings if finding.rule == 12 and finding.status == "manual_review"]
    assert {finding.question for finding in mismatches} == {1, 2, 3}


def test_independent_processing_sections_change_only_selected_content(tmp_path: Path):
    picture = tmp_path / "diagram.png"
    Image.new("RGB", (40, 30), "white").save(picture)

    tagged_source = tmp_path / "tagged_sections.docx"
    tagged = Document()
    for text in [
        "@Question: 1@", "@Type: FIB@", "@Question id: project13_q1@", "@New snippet id: 13001@",
        "@Difficulty level: Easy@", "@Objective: Application@", "@Question:@",
        "Find the roots of the equation x²−25=0.",
    ]:
        tagged.add_paragraph(text)
    tagged.add_picture(str(picture))
    for text in ["@e@", "@Answers:@", "5, -5", "@e@", "@Solution:@", "@e@"]:
        tagged.add_paragraph(text)
    tagged.save(tagged_source)

    image_only = process_docx(
        tagged_source,
        tagged_source.name,
        ProcessorOptions(
            processing_mode="custom",
            add_cms_tags=False,
            process_images=True,
            format_math_structure=False,
            project_question_prefix="fallback_q",
        ),
        tmp_path / "image_only",
    )
    image_xml = etree.fromstring(ZipFile(image_only.output_path).read("word/document.xml"))
    assert not image_xml.xpath(".//m:oMath", namespaces={"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"})
    assert image_xml.xpath(
        ".//wp:docPr/@descr",
        namespaces={"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"},
    ) == ["project13_q1_1.gif"]

    math_only = process_docx(
        tagged_source,
        tagged_source.name,
        ProcessorOptions(
            processing_mode="custom",
            add_cms_tags=False,
            process_images=False,
            format_math_structure=True,
        ),
        tmp_path / "math_only",
    )
    math_xml = etree.fromstring(ZipFile(math_only.output_path).read("word/document.xml"))
    assert math_xml.xpath(".//m:oMath", namespaces={"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"})
    assert math_only.image_count == 0
    assert not math_xml.xpath(
        ".//wp:docPr/@descr",
        namespaces={"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"},
    )

    plain_source = tmp_path / "plain_tags.docx"
    plain = Document()
    for text in ["Question 1: Easy, Application", "Question type: FIB", "What is 1/4 of 20?", "Answer: 5"]:
        plain.add_paragraph(text)
    plain.save(plain_source)
    tags_only = process_docx(
        plain_source,
        plain_source.name,
        ProcessorOptions(
            processing_mode="custom",
            add_cms_tags=True,
            process_images=False,
            format_math_structure=False,
            project_question_prefix="project13_q",
            start_snippet_id=13001,
        ),
        tmp_path / "tags_only",
    )
    tags_doc = Document(tags_only.output_path)
    tags_text = "\n".join(paragraph.text for paragraph in tags_doc.paragraphs)
    assert "@Question: 1@" in tags_text and "@Question id: project13_q1 @" in tags_text
    assert "1/4" in tags_text
    assert not tags_doc.element.xpath(".//m:oMath")
    report = __import__("json").loads(tags_only.report_path.read_text(encoding="utf-8"))
    assert report["selected_processing_sections"] == {
        "add_cms_tags": True,
        "process_images": False,
        "format_math_structure": False,
    }


def test_legacy_vml_picture_is_extracted_and_receives_alt_text(tmp_path: Path):
    source = tmp_path / "legacy_vml.docx"
    picture = tmp_path / "legacy.png"
    Image.new("RGB", (48, 32), "white").save(picture)
    doc = Document()
    for text in [
        "@Question: 1@",
        "@Type: FIB@",
        "@Question id: projectvml_q1@",
        "@New snippet id: 1@",
        "@Question:@",
        "Use the legacy diagram.",
    ]:
        doc.add_paragraph(text)

    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    run.add_picture(str(picture))
    drawing = run._r.xpath("./w:drawing")[0]
    relationship_id = drawing.xpath(".//a:blip")[0].get(qn("r:embed"))
    run._r.remove(drawing)
    word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    vml_ns = "urn:schemas-microsoft-com:vml"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pict = etree.Element(f"{{{word_ns}}}pict", nsmap={"w": word_ns, "v": vml_ns, "r": rel_ns})
    shape = etree.SubElement(pict, f"{{{vml_ns}}}shape")
    shape.set("alt", "Legacy diagram")
    image_data = etree.SubElement(shape, f"{{{vml_ns}}}imagedata")
    image_data.set(f"{{{rel_ns}}}id", relationship_id)
    run._r.append(pict)

    for text in ["@e@", "@Answers:@", "1", "@e@", "@Solution:@", "One.", "@e@"]:
        doc.add_paragraph(text)
    doc.save(source)

    result = process_docx(
        source,
        source.name,
        ProcessorOptions(processing_mode="images_only", project_question_prefix="fallback_q"),
        tmp_path / "storage",
    )
    xml = etree.fromstring(ZipFile(result.output_path).read("word/document.xml"))
    alt_values = xml.xpath(".//v:shape/@alt", namespaces={"v": vml_ns})

    assert result.image_count == 1
    assert alt_values == ["projectvml_q1_1.gif"]
    with ZipFile(result.images_zip_path) as archive:
        assert [name for name in archive.namelist() if name.lower().endswith(".gif")] == ["projectvml_q1_1.gif"]
